from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = PROJECT_ROOT / "docs"
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table):
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), LIGHT_GRAY)


def configure_document(doc, subtitle):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.03

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style.font.size = Pt(26)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style_properties = title_style._element.get_or_add_pPr()
    title_style_border = title_style_properties.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_properties.remove(title_style_border)

    for name, size in (("Heading 1", 16), ("Heading 2", 12.5)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(doc.core_properties.title)
    title_properties = title._p.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = intro.add_run(subtitle)
    run.italic = True
    run.font.size = Pt(11)
    intro.paragraph_format.space_after = Pt(16)


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_module_table(doc):
    rows = [
        ("Interface", "index.html and style.css", "Camera view, status, calibration control, and session metrics"),
        ("Tracking", "script.js and MediaPipe Pose", "Estimates landmarks and draws the live skeleton"),
        ("Validation", "isRealHumanUpperBody", "Rejects frames without plausible face and shoulder geometry"),
        ("Classification", "predictPosture", "Compares normalized geometry with the active baseline"),
        ("Reference data", "data/hf_posture_dataset.csv", "Supports exploration but does not train the live classifier"),
        ("Documentation", "docs and tools/generate_docs.py", "Keeps project reports reproducible"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = (Inches(1.2), Inches(2.05), Inches(3.65))
    headers = ("Module", "Files or function", "Responsibility")
    for index, (cell, text, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = width
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for cell, text, width in zip(cells, row_data, widths):
            cell.width = width
            cell.text = text
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, PALE_BLUE)
    set_table_borders(table)


def create_conference_report():
    doc = Document()
    doc.core_properties.title = "Posture Monitor Conference Review Report"
    doc.core_properties.subject = "Implementation and architecture review"
    configure_document(
        doc,
        "Architecture implementation evidence limitations and demonstration checklist",
    )

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "Posture Monitor is a working static web application that performs pose estimation and "
        "posture classification in the browser. The current implementation uses MediaPipe Pose, "
        "normalized face and shoulder geometry, and user calibration. It does not upload camera "
        "frames or require an application backend. The classifier is heuristic and should be "
        "presented as a posture aid rather than a trained medical or diagnostic model."
    )

    doc.add_heading("System Architecture", level=1)
    doc.add_paragraph(
        "Camera frame  >  MediaPipe Pose  >  upper body validation  >  f1 and f2 feature "
        "calculation  >  calibrated threshold classification  >  status counters and canvas"
    )
    doc.add_paragraph(
        "The browser owns the complete runtime path. It requests one camera stream, displays a "
        "preview while MediaPipe initializes, and then sends frames through a sequential "
        "request-animation-frame loop. This avoids duplicate camera streams and overlapping pose "
        "requests."
    )

    doc.add_heading("Implementation Modules", level=1)
    add_module_table(doc)

    doc.add_heading("Geometric Classification", level=1)
    doc.add_paragraph(
        "The application derives two scale-independent ratios. Feature f1 is the absolute "
        "difference between the left and right face-to-shoulder distances, multiplied by ten and "
        "divided by shoulder width. Feature f2 is the sum of those face-to-shoulder distances "
        "divided by shoulder width."
    )
    doc.add_paragraph(
        "A calibration action stores the current f1 and f2 values as the user's upright baseline. "
        "The classifier reports slouching when f2 falls sufficiently below that baseline and "
        "leaning when f1 moves sufficiently far from it. Calibration remains disabled until a "
        "valid upper body measurement is available."
    )

    validation_heading = doc.add_heading("Human Validation", level=1)
    validation_heading.paragraph_format.page_break_before = True
    add_bullets(
        doc,
        [
            "Requires strong visibility for the nose eyes and shoulders.",
            "Requires at least one visible ear.",
            "Checks that the shoulders appear below the nose.",
            "Rejects implausibly narrow shoulder spans.",
            "Checks the head-to-shoulder distance ratio.",
        ],
    )

    doc.add_heading("Reference Dataset", level=1)
    doc.add_paragraph(
        "The included ConfiDetect CSV contains 5,949 data records and uses "
        "Upright, Stiff, and Slouched posture labels. It does not contain a lateral-lean class, "
        "and its columns do not reproduce the live f1 and f2 features. It is retained for "
        "exploration and reproducibility but is not claimed as training data for the browser "
        "classifier."
    )

    doc.add_heading("Demonstration Checklist", level=1)
    add_bullets(
        doc,
        [
            "Serve the project through localhost or HTTPS and allow camera access.",
            "Show the preview and wait for the upper body validation message to clear.",
            "Sit upright and select Calibrate posture.",
            "Demonstrate upright slouch and lateral lean states.",
            "Show that counters increment once per transition rather than once per frame.",
            "Explain that session values reset on page reload.",
        ],
    )

    doc.add_heading("Known Limitations", level=1)
    add_bullets(
        doc,
        [
            "Classification quality depends on lighting camera placement clothing and visibility.",
            "The default thresholds require user testing and are not clinical thresholds.",
            "MediaPipe assets and fonts are loaded from third-party content delivery networks.",
            "No long-term history persistence notification system or multi-user profile is included.",
        ],
    )

    doc.save(DOCS_DIR / "CONFERENCE_REPORT.docx")


def create_project_thesis():
    doc = Document()
    doc.core_properties.title = "Posture Monitor Engineering Report"
    doc.core_properties.subject = "Technical project report"
    configure_document(
        doc,
        "Design implementation validation and limitations of a browser posture monitor",
    )

    doc.add_heading("Project Overview", level=1)
    doc.add_paragraph(
        "Posture Monitor demonstrates how a browser can combine webcam access, pose estimation, "
        "geometric feature extraction, and immediate visual feedback without an application "
        "server. The result is a small deployable system with a transparent classification path. "
        "Its main engineering contribution is the integration of MediaPipe landmarks with "
        "calibrated normalized measurements and a guarded real-time processing loop."
    )

    doc.add_heading("Problem Definition", level=1)
    doc.add_paragraph(
        "A useful desktop posture aid must respond quickly, adapt to camera placement, and avoid "
        "sending sensitive video to a custom backend. A single fixed neck angle is fragile because "
        "body proportions, seating position, lens height, and screen tilt vary. The project "
        "therefore uses ratios normalized by shoulder width and lets the user establish an upright "
        "baseline."
    )

    doc.add_heading("Architecture", level=1)
    doc.add_paragraph(
        "The application is delivered as HTML CSS and JavaScript. After the user grants camera "
        "permission, one MediaStream supplies the hidden video element. MediaPipe Pose processes "
        "frames sequentially. The result callback validates landmark geometry, calculates the two "
        "posture features, assigns a state, and renders the camera frame and skeleton on a canvas."
    )
    doc.add_heading("Runtime Sequence", level=2)
    add_bullets(
        doc,
        [
            "Check that browser camera APIs are available.",
            "Request one user-facing video stream without audio.",
            "Render a temporary raw preview while MediaPipe initializes.",
            "Process frames sequentially to prevent concurrent pose calls.",
            "Validate the detected face and shoulders before classification.",
            "Update the status skeleton and transition counters.",
        ],
    )

    doc.add_heading("Feature Engineering", level=1)
    doc.add_paragraph(
        "For each valid frame, the system calculates a weighted incenter for the left and right "
        "face triangles formed by an eye an ear and the nose. It then measures each face center to "
        "its corresponding shoulder and measures the shoulder span."
    )
    doc.add_paragraph(
        "Feature f1 equals abs((s1 minus s2) times 10) divided by s3. Feature f2 equals "
        "(s1 plus s2) divided by s3. Dividing by shoulder span makes the features less sensitive "
        "to camera distance than raw pixel measurements."
    )

    classification_heading = doc.add_heading("Classification and Calibration", level=1)
    classification_heading.paragraph_format.page_break_before = True
    doc.add_paragraph(
        "The production classifier uses deterministic thresholds. Slouching has priority when f2 "
        "falls below the baseline by more than 0.12. Lateral leaning is reported when f1 differs "
        "from its baseline by more than 0.55. Otherwise the state is upright. These values are "
        "engineering defaults and require evaluation across a representative user sample."
    )
    doc.add_paragraph(
        "The calibration button remains disabled until a valid measurement exists. Selecting it "
        "stores the current values as the baseline for the session. The baseline and counters are "
        "held in memory and are intentionally cleared by a reload."
    )

    doc.add_heading("Dataset Assessment", level=1)
    doc.add_paragraph(
        "The repository includes a downloaded posture dataset for exploratory analysis. Its labels "
        "are Upright, Stiff, and Slouched and its recorded features differ from the browser's f1 "
        "and f2 definitions. Because it has no lateral-lean label and no direct feature parity, the "
        "current system does not train or validate its live classifier from that CSV. This "
        "distinction prevents an unsupported accuracy claim."
    )

    doc.add_heading("Interface and Accessibility", level=1)
    doc.add_paragraph(
        "The interface provides a responsive camera card, status overlay, calibration control, "
        "session timer, and two event counters. Status changes use both text and color. The status "
        "region is announced through an ARIA live region, the canvas has an accessible label, and "
        "the button exposes a disabled state before calibration is possible."
    )

    doc.add_heading("Privacy and Security", level=1)
    doc.add_paragraph(
        "The project code does not upload or persist camera frames. It does load MediaPipe and font "
        "assets from public CDNs, so an internet connection is required on first load and those "
        "providers can observe ordinary asset requests. Camera access also requires localhost or "
        "HTTPS under normal browser security rules."
    )

    doc.add_heading("Testing Strategy", level=1)
    add_bullets(
        doc,
        [
            "Static syntax checks for JavaScript and Python files.",
            "Markup checks for required element identifiers and script dependencies.",
            "Manual camera permission denial and secure-context tests.",
            "Manual tests for no person partial person upright slouch and lean states.",
            "Transition tests confirming counters do not increment on every frame.",
            "Responsive layout checks at desktop and narrow mobile widths.",
        ],
    )

    doc.add_heading("Limitations and Future Work", level=1)
    add_bullets(
        doc,
        [
            "Collect labeled f1 and f2 sequences across diverse users and camera positions.",
            "Separate sustained poor posture from brief natural movement with time-based smoothing.",
            "Add opt-in local persistence for session summaries.",
            "Bundle or self-host runtime assets for offline operation.",
            "Evaluate thresholds with a documented protocol before reporting accuracy.",
        ],
    )

    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph(
        "The project meets its core goal of private real-time posture feedback in a static web "
        "application. Its implementation is compact and explainable, and calibration addresses an "
        "important source of variation. The next technical priority is a labeled evaluation set "
        "whose features exactly match the production measurements."
    )

    doc.save(DOCS_DIR / "PROJECT_THESIS.docx")


def main():
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    create_conference_report()
    create_project_thesis()
    print(f"Updated Word reports in {DOCS_DIR}")


if __name__ == "__main__":
    main()
